"""
Document text extraction.

Each function takes raw file bytes and returns extracted plain text.
Keeping these pure functions (no I/O beyond the bytes given) makes
them easy to unit test and to reuse for MCP-fetched documents later
(Phase 4).
"""
import io

from pypdf import PdfReader
from docx import Document as DocxDocument
import openpyxl


class UnsupportedFileType(Exception):
    pass


def extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def extract_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Tables often carry the important facts in registration/legal docs
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                paragraphs.append(" | ".join(cells))
    return "\n".join(paragraphs)


def extract_txt(data: bytes) -> str:
    return data.decode("utf-8", errors="ignore")


def extract_xlsx(data: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


EXTRACTORS = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "txt": extract_txt,
    "xlsx": extract_xlsx,
}


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    extractor = EXTRACTORS.get(ext)
    if extractor is None:
        raise UnsupportedFileType(
            f"'.{ext}' is not supported. Supported types: "
            f"{', '.join(EXTRACTORS.keys())}"
        )
    text = extractor(data)
    if not text.strip():
        raise ValueError(
            f"No extractable text found in '{filename}' "
            "(it may be a scanned/image-only file)."
        )
    return text
